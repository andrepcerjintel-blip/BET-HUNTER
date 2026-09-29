"""Exportação CSV / XLSX / JSON e listas para copiar."""
import csv
import io
import json

from . import queries
from .db import jl
from .util import NAO_IDENTIFICADO, split_iso

SIMPLE_COLS = ["username", "url_perfil", "url_video", "evidencia", "dominio", "url_externa", "score", "status",
               "data_coleta", "observacao_analista"]
MISSION_COLS = ["NUMERO", "USERNAME", "URL_PERFIL", "URL_VIDEO", "DESCRICAO_EVIDENCIA", "TEXTO_EVIDENCIA", "PLATAFORMA",
                "DOMINIO", "URL_EXTERNA", "CODIGO_AFILIADO", "DATA_COLETA", "OBSERVACAO_ANALISTA"]
FULL_COLS = ["USERNAME", "URL_PERFIL", "NOME_EXIBIDO", "BIO", "URL_VIDEO", "DATA_COLETA", "HORA_COLETA", "FUSO_HORARIO",
             "LEGENDA", "TEXTO_RELEVANTE", "HASHTAGS", "LINK_BIO", "URL_ORIGINAL", "URL_INTERMEDIARIA", "URL_FINAL",
             "DOMINIO_FINAL", "PARAMETROS_URL", "PLATAFORMA_MENCIONADA", "JOGO_MENCIONADO", "CODIGO_PROMOCIONAL",
             "AFFILIATE_ID", "SCORE", "CLASSIFICACAO", "TIPO", "MOTIVO_SCORE", "EVIDENCIA", "FONTE_DESCOBERTA",
             "STATUS_VALIDACAO", "VISUAL_ANALYSIS", "PADRAO_PROMOCIONAL_RECORRENTE", "OBSERVACAO_ANALISTA", "QTD_EVIDENCIAS",
             "FONTE_CONSULTADA", "URL_FONTE_CONSULTADA"]


def _v(x):
    if x is None or x == "" or x == [] or x == "[]":
        return NAO_IDENTIFICADO
    return x


def _join(lst):
    return " | ".join(str(i) for i in lst) if lst else ""


def _select(conn, f, scope):
    f = dict(f)
    if scope == "confirmed":
        f["status"] = ["CONFIRMADO", "JÁ ENCAMINHADO"]
        f["view"] = "all"
    elif scope == "all":
        f["view"] = f.get("view", "all")
    where, args = queries.build_where(f)
    return conn.execute(f"SELECT id FROM candidates WHERE {where} ORDER BY {queries.ORDERS['priority']}", args).fetchall()


def build_rows(conn, f, scope="confirmed", kind="simple"):
    out = []
    for r in _select(conn, f, scope):
        d = queries.candidate_detail(conn, r["id"])
        evs = [e for e in d["evidences"] if e["kind"] != "relacao"]
        first = evs[0] if evs else {}
        video_ev = next((e for e in evs if e["url_video"]), {})
        best_link = next((l for l in d["links"] if l["url_final"] == d["link_final"] and l["url_original"] == d["link_original"]),
                         d["links"][0] if d["links"] else {})
        date, hora, fuso = split_iso(first.get("collected_at") or d["last_analyzed"])
        params = _join([f"{p['param']}={p['value']} ({p['type']}, {p['dominio']})" for l in d["links"] for p in l["params"]])
        motivo = "; ".join(f"{'+' if x['pts'] > 0 else ''}{x['pts']} {x['label']}" + (f" [{x['detail']}]" if x["detail"] else "")
                           for x in d["reasons"])
        if kind == "mission":
            ev_txt = video_ev or first
            out.append({
                "NUMERO": len(out) + 1, "USERNAME": d["username"], "URL_PERFIL": d["profile_url"],
                "URL_VIDEO": _v(d["video_url"]), "DESCRICAO_EVIDENCIA": _v(d["main_evidence"]),
                "TEXTO_EVIDENCIA": _v((ev_txt.get("caption") or ev_txt.get("text") or d["bio"] or "")),
                "PLATAFORMA": _v(_join(d["platforms"])), "DOMINIO": _v(d["domain_final"]), "URL_EXTERNA": _v(d["link_final"]),
                "CODIGO_AFILIADO": _v(_join(d["affiliate_ids"] + d["codes"])),
                "DATA_COLETA": f"{date} {hora} {fuso}".strip() or NAO_IDENTIFICADO,
                "OBSERVACAO_ANALISTA": d["analyst_note"] or ""})
        elif kind == "simple":
            out.append({
                "username": d["username"], "url_perfil": d["profile_url"], "url_video": _v(d["video_url"]),
                "evidencia": _v(d["main_evidence"]), "dominio": _v(d["domain_final"]), "url_externa": _v(d["link_final"]),
                "score": d["score"], "status": d["status"], "data_coleta": f"{date} {hora} {fuso}".strip() or NAO_IDENTIFICADO,
                "observacao_analista": d["analyst_note"] or ""})
        else:
            out.append({
                "USERNAME": d["username"], "URL_PERFIL": d["profile_url"], "NOME_EXIBIDO": _v(d["display_name"]),
                "BIO": _v(d["bio"]), "URL_VIDEO": _v(d["video_url"]), "DATA_COLETA": date or NAO_IDENTIFICADO,
                "HORA_COLETA": hora or NAO_IDENTIFICADO, "FUSO_HORARIO": fuso or NAO_IDENTIFICADO,
                "LEGENDA": _v(video_ev.get("caption")), "TEXTO_RELEVANTE": _v(_join([e["caption"] or e["text"] for e in evs[:3]])),
                "HASHTAGS": _v(_join(["#" + h for h in d["hashtags"]])), "LINK_BIO": _v(d["bio_link"]),
                "URL_ORIGINAL": _v(d["link_original"]), "URL_INTERMEDIARIA": _v(_join(best_link.get("chain", []))),
                "URL_FINAL": _v(d["link_final"]), "DOMINIO_FINAL": _v(d["domain_final"]),
                "PARAMETROS_URL": _v(params), "PLATAFORMA_MENCIONADA": _v(_join(d["platforms"])),
                "JOGO_MENCIONADO": _v(_join(d["games"])), "CODIGO_PROMOCIONAL": _v(_join(d["codes"])),
                "AFFILIATE_ID": _v(_join(d["affiliate_ids"])), "SCORE": d["score"], "CLASSIFICACAO": d["classification"],
                "TIPO": d["content_type"], "MOTIVO_SCORE": _v(motivo), "EVIDENCIA": _v(d["main_evidence"]),
                "FONTE_DESCOBERTA": _v(_join(d["sources"])), "STATUS_VALIDACAO": d["status"],
                "VISUAL_ANALYSIS": d.get("visual_analysis") or "não disponível",
                "PADRAO_PROMOCIONAL_RECORRENTE": "SIM" if d["recurring"] else "NÃO",
                "OBSERVACAO_ANALISTA": d["analyst_note"] or "", "QTD_EVIDENCIAS": d["evidence_count"],
                "FONTE_CONSULTADA": _v(first.get("source")), "URL_FONTE_CONSULTADA": _v(first.get("source_url"))})
    return out


def export(conn, f, *, fmt="csv", kind="simple", scope="confirmed"):
    """-> (bytes, mimetype, filename)"""
    from .util import now_iso
    stamp = now_iso()[:19].replace(":", "").replace("-", "").replace("T", "_")
    name = f"bethunter_{kind}_{scope}_{stamp}"
    if fmt == "json":
        rows = build_rows(conn, f, scope, "full")
        full = []
        for r in _select(conn, f, scope):
            d = queries.candidate_detail(conn, r["id"])
            full.append(d)
        payload = {"exportado_em": now_iso(), "escopo": scope, "quantidade": len(full), "resumo": rows, "candidatos": full}
        return (json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8"), "application/json",
                name + ".json")
    rows = build_rows(conn, f, scope, kind)
    cols = {"simple": SIMPLE_COLS, "mission": MISSION_COLS}.get(kind, FULL_COLS)
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "resultados"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r[c] for c in cols])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(len(c) + 2, 14), 50)
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", name + ".xlsx"
    sio = io.StringIO()
    w = csv.DictWriter(sio, fieldnames=cols)
    w.writeheader()
    w.writerows(rows)
    return ("﻿" + sio.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", name + ".csv"


def copy_list(conn, f, what="profiles", scope="confirmed"):
    rows = _select(conn, f, scope)
    out = []
    for r in rows:
        c = conn.execute("SELECT username, profile_url FROM candidates WHERE id=?", (r["id"],)).fetchone()
        if what == "profiles":
            out.append(c["profile_url"])
        elif what == "usernames":
            out.append("@" + c["username"])
        elif what == "videos":
            for e in conn.execute("SELECT DISTINCT url_video FROM evidences WHERE candidate_id=? AND url_video!='' AND kind!='relacao'", (r["id"],)):
                out.append(e[0])
    return "\n".join(dict.fromkeys(out))
